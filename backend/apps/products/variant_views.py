"""API for the Add Multiple Items / Product Variant Master screen.

  GET    /api/products/variant-groups/?search=        parent products (with variant counts)
  GET    /api/products/variant-groups/{id}/           one parent with every variant + stock
  GET    /api/products/variant-groups/for-product/?product=<id>
  POST   /api/products/variant-groups/bulk-save/      validate + save (dry_run: true = validate only)
  GET    /api/products/variant-groups/import-template/   sample .xlsx
  POST   /api/products/variant-groups/import-preview/    parse + validate an uploaded sheet (writes nothing)
  POST   /api/products/variant-groups/import-commit/     save the valid rows of a previewed sheet
  GET    /api/products/variant-groups/export/?group=     .xlsx of variants (one parent, or all)
  CRUD   /api/products/variant-templates/

Permissions follow the rest of the products API (see ProductViewSet).
"""
import io
import os
import tempfile
from collections import OrderedDict

from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.serializers import ModelSerializer, ValidationError

from apps.common.branch_context import get_active_branch

from .import_engine import detect_ext, field_getter, open_row_stream, resolve_field_headers
from .models import Product, ProductVariantGroup, ProductVariantTemplate
from .variant_service import (
    CONTEXT_INVENTORY, MAX_ROWS_PER_REQUEST, VariantValidationError, save_variants,
    serialize_group, serialize_variant,
)

IMPORT_MAX_BYTES = 15 * 1024 * 1024

# Sheet columns, in template order: (field, header, aliases)
IMPORT_COLUMNS = [
    ('product_name', 'Product Name', ['parent product', 'item name', 'product']),
    ('name', 'Variant Name', ['variant', 'long name', 'variant name']),
    ('base_sku', 'Base SKU', ['parent sku', 'style code']),
    ('sku', 'SKU', ['item code', 'product code', 'sku code']),
    ('barcode', 'Barcode', ['ean', 'upc', 'bar code']),
    ('category', 'Category', ['product category']),
    ('sub_category', 'Subcategory', ['sub category', 'sub-category']),
    ('group', 'Group', ['product group']),
    ('brand', 'Brand', ['manufacturer']),
    ('model_number', 'Model Number', ['model', 'model no']),
    ('color', 'Colour', ['color', 'colour name']),
    ('size', 'Size', ['frame size']),
    ('power', 'Power', ['sph', 'specification', 'spec']),
    ('hsn_code', 'HSN Code', ['hsn', 'hsn/sac']),
    ('unit', 'Unit', ['uom']),
    ('cost_price', 'Cost Price', ['purchase price', 'cost']),
    ('mrp', 'MRP', ['max retail price']),
    ('sales_price', 'Sales Price', ['selling price', 'sale price', 'retail price']),
    ('tax_percent', 'Tax %', ['gst', 'gst %', 'tax', 'tax rate']),
    ('opening_stock', 'Opening Qty', ['opening stock', 'qty', 'quantity', 'stock']),
]
IMPORT_ALIASES = OrderedDict((f, [h.lower()] + a) for f, h, a in IMPORT_COLUMNS)
VARIANT_ATTR_FIELDS = ('color', 'size', 'power')
GROUP_FIELDS = ('base_sku', 'category', 'sub_category', 'group', 'brand', 'model_number', 'hsn_code', 'unit', 'tax_percent')


class VariantTemplateSerializer(ModelSerializer):
    class Meta:
        model = ProductVariantTemplate
        fields = ['id', 'name', 'category_name', 'config', 'created_at', 'updated_at']

    # Templates are reusable configuration only — reject anything transactional.
    FORBIDDEN_KEYS = {'supplier', 'supplier_id', 'invoice_number', 'supplier_invoice_number', 'purchase_id', 'rows', 'variants'}

    def validate_config(self, value):
        if not isinstance(value, dict):
            raise ValidationError('Template config must be an object.')
        bad = self.FORBIDDEN_KEYS & set(value)
        if bad:
            raise ValidationError(f'Templates cannot store transaction data ({", ".join(sorted(bad))}).')
        return value


class ProductVariantTemplateViewSet(viewsets.ModelViewSet):
    queryset = ProductVariantTemplate.objects.all()
    serializer_class = VariantTemplateSerializer
    authentication_classes = []
    permission_classes = [permissions.AllowAny]
    pagination_class = None

    def destroy(self, request, *args, **kwargs):
        self.get_object().hard_delete()  # names are unique; a soft-deleted row would block reuse
        return Response(status=status.HTTP_204_NO_CONTENT)

    @action(detail=True, methods=['post'])
    def duplicate(self, request, pk=None):
        src = self.get_object()
        base = (request.data.get('name') or f'{src.name} (copy)').strip()[:150]
        name, n = base, 1
        while ProductVariantTemplate.all_objects.filter(name__iexact=name).exists():
            n += 1
            name = f'{base} {n}'[:150]
        copy = ProductVariantTemplate.objects.create(name=name, category_name=src.category_name, config=src.config)
        return Response(self.get_serializer(copy).data, status=status.HTTP_201_CREATED)


def _error_response(exc):
    return Response(exc.as_dict(), status=exc.status)


def _changed_by(request):
    user = getattr(request, 'user', None)
    if getattr(user, 'is_authenticated', False):
        return user.get_username()
    return (getattr(request, 'data', None) or {}).get('changed_by') or 'Administrator'


class ProductVariantGroupViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = ProductVariantGroup.objects.all()
    authentication_classes = []
    permission_classes = [permissions.AllowAny]
    pagination_class = None

    def list(self, request):
        qs = ProductVariantGroup.objects.all()
        search = (request.query_params.get('search') or '').strip()
        if search:
            # Filter by id first: filtering on `variants__…` after annotating would join the
            # variants a second time and multiply the counts below.
            ids = ProductVariantGroup.objects.filter(
                Q(name__icontains=search) | Q(base_sku__icontains=search) | Q(model_number__icontains=search)
                | Q(brand_name__icontains=search) | Q(variants__sku__icontains=search)
                | Q(variants__barcode__icontains=search)
            ).values('id')
            qs = qs.filter(id__in=ids)
        qs = qs.annotate(
            variant_count=Count('variants', filter=Q(variants__is_deleted=False)),
            total_stock=Sum('variants__stock', filter=Q(variants__is_deleted=False)),
        ).select_related('unit').order_by('-updated_at')
        try:
            limit = max(1, min(int(request.query_params.get('limit') or 50), 200))
        except ValueError:
            limit = 50
        out = []
        for g in qs[:limit]:
            row = serialize_group(g)
            row['variant_count'] = g.variant_count
            row['total_stock'] = g.total_stock or 0
            out.append(row)
        return Response(out)

    def retrieve(self, request, pk=None):
        return Response(serialize_group(self.get_object(), with_variants=True))

    @action(detail=False, methods=['get'], url_path='for-product')
    def for_product(self, request):
        """Manage Variants on any product: its group if it has one, else the product alone."""
        product = Product.objects.select_related('variant_group', 'tax').filter(pk=request.query_params.get('product')).first()
        if product is None:
            return Response({'detail': 'Product not found.'}, status=status.HTTP_404_NOT_FOUND)
        if product.variant_group_id:
            return Response(serialize_group(product.variant_group, with_variants=True))
        return Response({'id': None, 'standalone': True, 'variants': [serialize_variant(product)]})

    @action(detail=False, methods=['post'], url_path='bulk-save')
    def bulk_save(self, request):
        try:
            branch = get_active_branch(request)
        except Exception:
            branch = None
        try:
            result = save_variants(request.data, user=request.user, branch=branch, changed_by=_changed_by(request))
        except VariantValidationError as exc:
            return _error_response(exc)
        code = status.HTTP_200_OK if result.get('dry_run') or result.get('replayed') or not result.get('created') else status.HTTP_201_CREATED
        return Response(result, status=code)

    # ------------------------------------------------------------------ Excel

    @action(detail=False, methods=['get'], url_path='import-template')
    def import_template(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Variants'
        ws.append([h for _f, h, _a in IMPORT_COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1E3A8A')
        ws.append(['Example Frame', '', 'EXFRM', '', '', 'Frames', 'Full Rim', 'Optical Frame', '', 'EX-100',
                   'Black', 'M', '', '90031100', 'PCS', 500, 1500, 1200, 12, 0])
        ws.append(['Example Frame', '', 'EXFRM', '', '', 'Frames', 'Full Rim', 'Optical Frame', '', 'EX-100',
                   'Brown', 'L', '', '90031100', 'PCS', 500, 1500, 1200, 12, 0])
        for i, (_f, h, _a) in enumerate(IMPORT_COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(h) + 4)
        notes = wb.create_sheet('Instructions')
        for line in [
            'One row per variant. Rows with the same Product Name become variants of one product.',
            'Colour / Size / Power are the variant attributes; leave them blank for a single item.',
            'SKU and Barcode are optional: blank SKUs are generated from Base SKU + attributes.',
            'Barcode must not already belong to another product. Delete the two example rows before importing.',
            'Opening Qty is only used when importing from Inventory → Products.',
        ]:
            notes.append([line])
        notes.column_dimensions['A'].width = 110
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="add-multiple-items-template.xlsx"'
        return resp

    @action(detail=False, methods=['post'], url_path='import-preview')
    def import_preview(self, request):
        upload = request.FILES.get('file')
        if not upload:
            return Response({'detail': 'Choose a file to import.'}, status=status.HTTP_400_BAD_REQUEST)
        ext = detect_ext(upload.name)
        if ext is None:
            return Response({'detail': 'Unsupported file. Use .xlsx, .xls or .csv.'}, status=status.HTTP_400_BAD_REQUEST)
        if upload.size > IMPORT_MAX_BYTES:
            return Response({'detail': 'File is too large for a preview import (max 15 MB). Use Excel Import Manager for large catalogues.'}, status=status.HTTP_400_BAD_REQUEST)
        context = request.data.get('context') or CONTEXT_INVENTORY

        fd, path = tempfile.mkstemp(suffix=f'.{ext}')
        try:
            with os.fdopen(fd, 'wb') as fh:
                for chunk in upload.chunks():
                    fh.write(chunk)
            try:
                headers, stream = open_row_stream(path, ext)
                resolved = resolve_field_headers(headers, IMPORT_ALIASES)
                raw_rows = []
                for row_number, row in enumerate(stream, start=2):
                    get = field_getter(resolved, row)
                    values = {f: get(f) for f in IMPORT_ALIASES}
                    if not any(values.values()):
                        continue
                    values['row_number'] = row_number
                    raw_rows.append(values)
                    if len(raw_rows) > MAX_ROWS_PER_REQUEST:
                        return Response({'detail': f'More than {MAX_ROWS_PER_REQUEST} rows. Split the file, or use Excel Import Manager for large catalogues.'}, status=status.HTTP_400_BAD_REQUEST)
            except ImportError:
                return Response({'detail': 'Old .xls files are not supported on this server. Save the sheet as .xlsx and try again.'}, status=status.HTTP_400_BAD_REQUEST)
            except Exception:
                return Response({'detail': 'Could not read the file. Check that it is a valid Excel/CSV sheet.'}, status=status.HTTP_400_BAD_REQUEST)
        finally:
            try:
                os.remove(path)
            except OSError:
                pass

        if 'product_name' not in resolved and 'name' not in resolved:
            return Response({'detail': 'No "Product Name" column found. Download the sample template for the expected columns.',
                             'mapping': resolved}, status=status.HTTP_400_BAD_REQUEST)

        groups = _groups_from_rows(raw_rows, context)
        for g in groups:
            _annotate_errors(g)
        return Response({
            'mapping': {f: resolved.get(f) for f in IMPORT_ALIASES},
            'unmapped_headers': [h for h in headers if h and h not in resolved.values()],
            'groups': groups,
            'total_rows': len(raw_rows),
            'valid_rows': sum(1 for g in groups for v in g['variants'] if not v['errors']),
            'error_rows': sum(1 for g in groups for v in g['variants'] if v['errors']),
        })

    @action(detail=False, methods=['post'], url_path='import-commit')
    def import_commit(self, request):
        """Save every valid row of a previewed sheet; each product is saved atomically on its
        own, and rows that fail validation are reported back instead of being written."""
        groups = request.data.get('groups') or []
        context = request.data.get('context') or CONTEXT_INVENTORY
        try:
            branch = get_active_branch(request)
        except Exception:
            branch = None
        imported, failed, errors, saved_groups = 0, 0, [], []
        for g in groups:
            payload = {
                'context': context, 'group': g.get('group') or {}, 'auto_sku': True,
                'auto_barcode': bool(request.data.get('auto_barcode')),
                'variants': [dict(v.get('data') or {}, row_number=v.get('row_number')) for v in g.get('variants') or []],
            }
            if not payload['variants']:
                continue
            # Re-validate server-side and drop the rows that fail; save the rest.
            try:
                save_variants({**payload, 'dry_run': True})
                good = payload['variants']
            except VariantValidationError as exc:
                bad_idx = {int(i) for i in exc.row_errors}
                for i, errs in exc.row_errors.items():
                    v = payload['variants'][int(i)]
                    for field, msg in errs.items():
                        errors.append({'row_number': v.get('row_number'), 'product': payload['group'].get('name'), 'field': field, 'message': msg})
                if exc.group_errors or exc.status != 400:
                    for field, msg in (exc.group_errors or {'_': exc.detail}).items():
                        errors.append({'row_number': None, 'product': payload['group'].get('name'), 'field': field, 'message': msg})
                    failed += len(payload['variants'])
                    continue
                failed += len(bad_idx)
                good = [v for i, v in enumerate(payload['variants']) if i not in bad_idx]
            if not good:
                continue
            try:
                result = save_variants({**payload, 'variants': good}, user=request.user, branch=branch, changed_by=_changed_by(request))
                imported += len(good)
                saved_groups.append(result.get('group') or {'name': payload['group'].get('name')})
            except VariantValidationError as exc:
                failed += len(good)
                errors.append({'row_number': None, 'product': payload['group'].get('name'), 'field': '_', 'message': exc.detail})
        return Response({'imported_rows': imported, 'failed_rows': failed, 'errors': errors, 'groups': saved_groups})

    @action(detail=False, methods=['get'])
    def export(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        qs = Product.objects.select_related('variant_group', 'tax').order_by('variant_group__name', 'name')
        group_id = request.query_params.get('group')
        qs = qs.filter(variant_group_id=group_id) if group_id else qs.filter(variant_group__isnull=False)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Variants'
        ws.append([h for _f, h, _a in IMPORT_COLUMNS] + ['Current Stock', 'Active'])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for p in qs.iterator(chunk_size=500):
            v = serialize_variant(p)
            g = p.variant_group
            attrs = v['attributes'] or {}
            row = {
                'product_name': g.name if g else p.name, 'name': v['name'], 'base_sku': g.base_sku if g else '',
                'sku': v['sku'], 'barcode': v['barcode'], 'category': v['category'], 'sub_category': v['sub_category'],
                'group': v['group'], 'brand': v['brand'], 'model_number': v['model_number'],
                'color': attrs.get('color', ''), 'size': attrs.get('size', ''), 'power': attrs.get('power', ''),
                'hsn_code': v['hsn_code'], 'unit': g.unit.name if g and g.unit_id else '',
                'cost_price': v['cost_price'], 'mrp': v['mrp'], 'sales_price': v['sales_price'],
                'tax_percent': v['tax_percent'], 'opening_stock': '',
            }
            ws.append([row[f] for f, _h, _a in IMPORT_COLUMNS] + [v['stock'], 'Yes' if v['is_active'] else 'No'])
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="product-variants.xlsx"'
        return resp


def _groups_from_rows(raw_rows, context):
    groups = OrderedDict()
    for r in raw_rows:
        parent = r.get('product_name') or r.get('name')
        key = parent.strip().lower()
        g = groups.get(key)
        if g is None:
            g = groups[key] = {'group': {'name': parent.strip()}, 'variants': []}
        for f in GROUP_FIELDS:
            if r.get(f) and not g['group'].get(f):
                g['group'][f] = r[f]
        attrs = {k: r[k] for k in VARIANT_ATTR_FIELDS if r.get(k)}
        data = {
            'name': r.get('name') if r.get('product_name') else '',
            'attributes': attrs,
            'sku': r.get('sku'), 'barcode': r.get('barcode'),
            'category': r.get('category'), 'sub_category': r.get('sub_category'), 'group': r.get('group'),
            'brand': r.get('brand'), 'model_number': r.get('model_number'), 'hsn_code': r.get('hsn_code'),
            'cost_price': r.get('cost_price'), 'mrp': r.get('mrp'), 'sales_price': r.get('sales_price'),
            'tax_percent': r.get('tax_percent'),
            'opening_stock': r.get('opening_stock') if context == CONTEXT_INVENTORY else '',
        }
        g['variants'].append({'row_number': r['row_number'], 'data': data, 'errors': {}})
    for g in groups.values():
        g['context'] = context
    return list(groups.values())


def _annotate_errors(g):
    try:
        res = save_variants({'context': g.get('context'), 'group': g['group'], 'dry_run': True,
                             'variants': [v['data'] for v in g['variants']]})
        for item in res.get('variants') or []:
            g['variants'][item['index']]['generated_sku'] = item['sku']
    except VariantValidationError as exc:
        for i, errs in exc.row_errors.items():
            g['variants'][int(i)]['errors'] = errs
        if exc.group_errors:
            g['group_errors'] = exc.group_errors
            for v in g['variants']:
                v['errors'] = v['errors'] or {'_group': '; '.join(exc.group_errors.values())}
